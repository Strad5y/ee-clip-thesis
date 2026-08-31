# class_table.py  – laedt die gefilterte Klassenliste und haengt Labels an POIs
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook

BASE_COLS = {"key", "value", "keep", "note"}  # alles andere gilt als Label-Spalte
QID_PATTERN = r"(Q\d+)"  # zum Bereinigen von Excel-Tippfehlern wie "Q1139861)"

def load_class_table(xlsx_path: str | Path, name: str = "Table_True") -> pd.DataFrame:
    """Liest 'Table_True' - ob Blatt oder Excel-Tabellenobjekt. Falls keins von beidem
    existiert, wird als Fallback das einzige Sheet der Datei gelesen (fuer eine bereits
    kuratierte, flache Klassenliste ohne Table_True-Struktur, z.B. osm_wikidata_keys.xlsx) -
    nur bei GENAU EINEM Sheet, sonst waere die Auswahl geraten statt explizit. Gibt nur
    keep=True zurueck, falls eine keep-Spalte vorhanden ist - sonst gelten alle Zeilen als
    aktiv (die Datei ist dann bereits die kuratierte Liste)."""
    xlsx_path = Path(xlsx_path)
    try:                                   # Fall A: es ist ein Arbeitsblatt
        df = pd.read_excel(xlsx_path, sheet_name=name)
    except ValueError:                     # Fall B: es ist ein benanntes Tabellenobjekt
        wb = load_workbook(xlsx_path, data_only=True)
        df = None
        for ws in wb.worksheets:
            if name in ws.tables:
                ref = ws.tables[name].ref            # z.B. "A1:E276"
                cells = ws[ref]
                header = [c.value for c in cells[0]]
                rows = [[c.value for c in row] for row in cells[1:]]
                df = pd.DataFrame(rows, columns=header)
                break
        if df is None:
            # Fall C: weder Blatt noch Tabellenobjekt "Table_True" gefunden
            if len(wb.sheetnames) != 1:
                raise ValueError(
                    f"'{name}' weder als Blatt noch als Tabelle gefunden, und {xlsx_path} hat "
                    f"{len(wb.sheetnames)} Sheets statt genau einem - Fallback nicht eindeutig."
                )
            df = pd.read_excel(xlsx_path, sheet_name=wb.sheetnames[0])

    df.columns = [str(c).strip() for c in df.columns]

    # keep robust normalisieren (Bool ODER Text 'WAHR'/'FALSCH' aus Excel-Handedits)
    if "keep" in df.columns:
        df["keep"] = df["keep"].astype(str).str.upper().isin(["TRUE", "WAHR", "1"])
        df = df[df["keep"]].copy()

    # key/value saeubern (fuer sauberen Merge)
    for c in ("key", "value"):
        df[c] = df[c].astype(str).str.strip()

    # wikidata saeubern: Excel enthaelt vereinzelt Tippfehler wie "Q1139861)" oder
    # fuehrende Leerzeichen - QID-Substring extrahieren, sonst schlagen spaetere
    # wikidata-Merges (Wikidata-API, build_poi_text.py) fuer diese Zeilen still fehl
    if "wikidata" in df.columns:
        cleaned = df["wikidata"].astype(str).str.strip().str.extract(QID_PATTERN, expand=False)
        n_bad = int((df["wikidata"].notna() & cleaned.isna()).sum())
        if n_bad > 0:
            print(f"  Warnung: {n_bad} 'wikidata'-Werte ohne gueltiges QID-Muster (Q\\d+), auf NaN gesetzt")
        df["wikidata"] = cleaned

    df = df.drop_duplicates(subset=["key", "value"]).reset_index(drop=True)
    return df


def attach_labels(poi_long: pd.DataFrame, class_df: pd.DataFrame) -> pd.DataFrame:
    """Inner-Join auf (key,value): filtert POI-Tags auf die behaltenen Klassen UND haengt deren Labels an."""
    label_cols = [c for c in class_df.columns if c not in BASE_COLS]  # feinlabel/hilucs_*/class_qid – was da ist
    keep_map = class_df[["key", "value", *label_cols]].copy()

    poi = poi_long.copy()
    poi["key"] = poi["key"].astype(str).str.strip()
    poi["value"] = poi["value"].astype(str).str.strip()

    merged = poi.merge(keep_map, on=["key", "value"], how="inner")  # inner => nicht-behaltene Tags fallen raus
    return merged


if __name__ == "__main__":
    cls = load_class_table("wikidata tags/osm_wikidata_keys.xlsx")
    print(f"{len(cls)} behaltene Klassen | Label-Spalten: "
          f"{[c for c in cls.columns if c not in BASE_COLS]}")
    print(cls.head())